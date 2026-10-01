"""
Build ``data/inputs/sutra_context.json`` — one record per Aṣṭādhyāyī sūtra —
from the sūtra-map workbook (base + ``Sutra overrides``), filling gaps from the
``~/ashtadhyayi`` checkout (ashtadhyayi.com export). One-way, idempotent and
deterministic: never hand-edit the output, re-run this script.

    python3 scripts/build_sutra_context.py [--workbook X.xlsx] [--ashtadhyayi DIR]

Conventions: unknown = ``null``, known-empty = ``[]``. Every record carries
``provenance[field] ∈ {workbook, override, ashtadhyayi.com}``. Disagreements
between the two sources are written to ``data/inputs/sutra_context.conflicts.json``
(and summarised in ``docs/SOURCE_CONFLICTS.md`` by hand) — never auto-resolved.
The XLOOKUP scheduling sheets are not read.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = Path("/Users/dr.ajayshukla/paanini-ashtadhyaayi-sutra-map (1).xlsx")
ASHTADHYAYI = Path("/Users/dr.ajayshukla/ashtadhyayi")
OUT = ROOT / "data/inputs/sutra_context.json"
OUT_SOURCE = ROOT / "data/inputs/sutra_context.SOURCE.json"
OUT_CONFLICTS = ROOT / "data/inputs/sutra_context.conflicts.json"

_DEV_DIGITS = str.maketrans("०१२३४५६७८९", "0123456789")
_VV = re.compile(r"^([0-9०-९])/([0-9०-९])$")
_ANU = re.compile(r"^(\d{5})\s*:\s*(.*)$")
_BRACKET_ID = re.compile(r"\[(\d+\.\d+\.\d+)\]")

# workbook sutra-sheet columns
C_KRAMA, C_KAUMUDI, C_ID, C_TYPE, C_TERM = 3, 5, 6, 7, 8
C_INFLUENCE, C_TEXT, C_PC, C_SC, C_ANU, C_PC_NOTES, C_TAGS = 11, 13, 14, 15, 16, 17, 18
# override-sheet columns → field
COMMENTARIES = ("kashika", "vasu_english", "vasu_english_summary")
COMMENTARY_PATH = "<ashtadhyayi>/<ref>/pada-<A>.<P>/<id>.md"
OVERRIDE_COLS = {1: "type", 2: "term", 3: "text", 4: "padaccheda_raw",
                 5: "samasaccheda", 6: "anuvritti_raw", 7: "pada_tags"}


def krama_to_id(k) -> str:
    s = str(int(float(k)))
    return f"{int(s[0])}.{int(s[1])}.{int(s[2:])}"


def id_key(sid: str) -> tuple[int, ...]:
    return tuple(int(p) for p in sid.split("."))


def _s(v) -> str | None:
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def parse_types(raw: str | None) -> list[str] | None:
    if raw is None:
        return None
    return [t.strip() for t in raw.split(";") if t.strip()]


def parse_padaccheda(raw: str | None) -> list[dict] | None:
    """'वृद्धिः १/१ आदैच् १/१' → [{word, vibhakti, vacana, note}]; '0/0' = avyaya."""
    if raw is None:
        return None
    out: list[dict] = []
    for tok in re.findall(r"\([^)]*\)|[^\s()]+", raw):
        m = _VV.match(tok)
        if m and out and out[-1]["vibhakti"] is None:
            out[-1]["vibhakti"] = int(m.group(1).translate(_DEV_DIGITS))
            out[-1]["vacana"] = int(m.group(2).translate(_DEV_DIGITS))
        elif tok.startswith("(") and out:
            out[-1]["note"] = tok[1:-1].strip()
        elif tok not in {"|", "।"}:
            out.append({"word": tok, "vibhakti": None, "vacana": None, "note": None})
    return out


def parse_anuvritti_workbook(raw: str | None) -> list[dict] | None:
    if raw is None:
        return None
    if raw.lower() == "na":
        return []
    out = []
    for part in raw.split("|"):
        m = _ANU.match(part.strip())
        if m:
            out.append({"word": m.group(2).strip(), "source_sutra": krama_to_id(m.group(1))})
    return out


def parse_range(raw: str | None) -> list[str] | None:
    if raw is None:
        return None
    m = re.match(r"^(\d{5})\s*-\s*(\d{5})$", raw)
    return [krama_to_id(m.group(1)), krama_to_id(m.group(2))] if m else None


def _read(path: Path) -> str | None:
    try:
        return path.read_text(encoding="utf-8").strip()
    except FileNotFoundError:
        return None


def _pada_path(base: Path, sub: str, sid: str, ext: str) -> Path:
    a, p, _ = sid.split(".")
    return base / sub / f"pada-{a}.{p}" / f"{sid}.{ext}"


def ash_heads(base: Path, sid: str) -> list[str] | None:
    raw = _read(_pada_path(base, "adhikara", sid, "txt"))
    if raw is None:
        return None
    return sorted({h for h in _BRACKET_ID.findall(raw)}, key=id_key)


def ash_anuvritti_sources(base: Path, sid: str) -> list[str] | None:
    raw = _read(_pada_path(base, "anuvritti", sid, "txt"))
    if raw is None:
        return None
    return sorted(set(_BRACKET_ID.findall(raw)), key=id_key)


def _norm_type(types: list[str] | None) -> frozenset[str]:
    return frozenset(re.split(r"\s+-\s+", t)[0] for t in (types or []))


def build(workbook: Path, ash: Path):
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    base_rows = [r for r in list(wb["sutra"].iter_rows(values_only=True))[1:] if r and r[C_KRAMA]]
    overrides = {krama_to_id(r[0]): r for r in list(wb["Sutra overrides"].iter_rows(values_only=True))[1:]
                 if r and r[0]}
    basics = {d["id"]: d for d in json.loads((ash / "sutraBasics.json").read_text(encoding="utf-8"))["sutraDetails"]}

    records: dict[str, dict] = {}
    extras: dict[str, dict] = {}
    conflicts = {"type": [], "adhikara_heads": [], "anuvritti": []}
    applied_overrides = 0

    for r in base_rows:
        sid = krama_to_id(r[C_KRAMA])
        rec = {
            "id": sid,
            "type": _s(r[C_TYPE]),
            "term": _s(r[C_TERM]),
            "text": _s(r[C_TEXT]),
            "padaccheda_raw": _s(r[C_PC]),
            "samasaccheda": _s(r[C_SC]),
            "anuvritti_raw": _s(r[C_ANU]),
            "padaccheda_notes": _s(r[C_PC_NOTES]),
            "pada_tags": _s(r[C_TAGS]),
            "adhikara_range": parse_range(_s(r[C_INFLUENCE])),
            "kaumudi_krama": int(r[C_KAUMUDI]) if r[C_KAUMUDI] else None,
        }
        prov = {k: "workbook" for k, v in rec.items() if k != "id" and v is not None}
        ov = overrides.get(sid)
        if ov is not None:
            applied_overrides += 1
            for col, field in OVERRIDE_COLS.items():
                v = _s(ov[col]) if col < len(ov) else None
                if v is not None:
                    rec[field] = v
                    prov[field] = "override"

        if sid not in basics:
            extras[sid] = {**rec, "provenance": prov}
            continue

        b = basics[sid]
        if rec["type"] is None and _s(b.get("प्रकारः")):
            rec["type"] = _s(b["प्रकारः"])
            prov["type"] = "ashtadhyayi.com"
        if rec["text"] is None and _s(b.get("सूत्रम्")):
            rec["text"] = _s(b["सूत्रम्"])
            prov["text"] = "ashtadhyayi.com"
        if rec["kaumudi_krama"] is None and b.get("कौमुदीक्रमसङ्ख्या") not in (None, ""):
            rec["kaumudi_krama"] = int(b["कौमुदीक्रमसङ्ख्या"])
            prov["kaumudi_krama"] = "ashtadhyayi.com"
        if rec["padaccheda_raw"] is None:
            pc = _read(_pada_path(ash, "padachcheda", sid, "txt"))
            if pc:
                rec["padaccheda_raw"] = pc
                prov["padaccheda_raw"] = "ashtadhyayi.com"

        rec["type"] = parse_types(rec["type"])
        rec["padaccheda"] = parse_padaccheda(rec["padaccheda_raw"])
        if rec["padaccheda"] is not None:
            prov["padaccheda"] = prov["padaccheda_raw"]
        rec["anuvritti"] = parse_anuvritti_workbook(rec["anuvritti_raw"])
        if rec["anuvritti"] is not None:
            prov["anuvritti"] = prov["anuvritti_raw"]

        heads = ash_heads(ash, sid)
        rec["adhikara_heads"] = heads
        if heads is not None:
            prov["adhikara_heads"] = "ashtadhyayi.com"
        topic = _read(_pada_path(ash, "topic", sid, "txt"))
        rec["topic"] = topic or None
        if topic:
            prov["topic"] = "ashtadhyayi.com"
        rec["commentary_refs"] = [k for k in COMMENTARIES if _pada_path(ash, k, sid, "md").exists()]
        rec["provenance"] = dict(sorted(prov.items()))
        records[sid] = rec

        # ── cross-source disagreements (logged, not resolved) ──
        bt = parse_types(_s(b.get("प्रकारः")))
        if prov.get("type") != "ashtadhyayi.com" and bt and _norm_type(rec["type"]) != _norm_type(bt):
            conflicts["type"].append({"id": sid, "workbook": rec["type"], "ashtadhyayi.com": bt})
        ash_anu = ash_anuvritti_sources(ash, sid)
        if ash_anu is not None and rec["anuvritti"] is not None:
            wb_src = sorted({a["source_sutra"] for a in rec["anuvritti"]}, key=id_key)
            if not set(ash_anu) <= set(wb_src):
                conflicts["anuvritti"].append({"id": sid, "workbook": wb_src, "ashtadhyayi.com": ash_anu})

    # adhikāra extent: workbook range on each head vs the sūtras the corpus says it governs
    governed: dict[str, list[str]] = {}
    for sid, rec in records.items():
        for h in rec["adhikara_heads"] or []:
            governed.setdefault(h, []).append(sid)
    for sid, rec in sorted(records.items(), key=lambda kv: id_key(kv[0])):
        rng = rec["adhikara_range"]
        span = sorted(governed.get(sid, []), key=id_key)
        corpus = [span[0], span[-1]] if span else None
        if rng is not None and corpus is not None and rng != corpus:
            conflicts["adhikara_heads"].append({"head": sid, "workbook_range": rng, "ashtadhyayi.com_extent": corpus})
        elif (rng is None) != (corpus is None):
            conflicts["adhikara_heads"].append({"head": sid, "workbook_range": rng, "ashtadhyayi.com_extent": corpus})

    ordered = {sid: records[sid] for sid in sorted(records, key=id_key)}
    return ordered, extras, conflicts, applied_overrides, len(overrides)


def _dump(obj) -> str:
    return json.dumps(obj, ensure_ascii=False, indent=1, sort_keys=True) + "\n"


def _dump_by_line(top: dict[str, dict]) -> str:
    """Valid JSON, one sūtra record per line (diff-friendly, ~half the indented size)."""
    parts = []
    for key, recs in top.items():
        body = ",\n".join(f"  {json.dumps(k, ensure_ascii=False)}: "
                          f"{json.dumps(v, ensure_ascii=False, sort_keys=True)}" for k, v in recs.items())
        parts.append(f" {json.dumps(key)}: {{\n{body}\n }}" if recs else f" {json.dumps(key)}: {{}}")
    return "{\n" + ",\n".join(parts) + "\n}\n"


def main(argv=None) -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", type=Path, default=WORKBOOK)
    ap.add_argument("--ashtadhyayi", type=Path, default=ASHTADHYAYI)
    ap.add_argument("--out", type=Path, default=OUT)
    a = ap.parse_args(argv)

    records, extras, conflicts, applied, n_over = build(a.workbook, a.ashtadhyayi)
    a.out.write_text(_dump_by_line({"extras": extras, "sutras": records}), encoding="utf-8")
    OUT_CONFLICTS.write_text(_dump({k: {"count": len(v), "items": v} for k, v in conflicts.items()}), encoding="utf-8")
    try:
        commit = subprocess.run(["git", "-C", str(a.ashtadhyayi), "rev-parse", "--short=10", "HEAD"],
                                capture_output=True, text=True, check=True).stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        commit = None
    OUT_SOURCE.write_text(_dump({
        "workbook": {"path": str(a.workbook), "sha256": hashlib.sha256(a.workbook.read_bytes()).hexdigest(),
                     "sheets_read": ["sutra", "Sutra overrides"], "overrides": n_over, "overrides_applied": applied},
        "ashtadhyayi": {"path": str(a.ashtadhyayi), "git_commit": commit,
                        "read": ["sutraBasics.json", "adhikara/", "anuvritti/", "padachcheda/", "topic/",
                                 "kashika/ (refs)", "vasu_english*/ (refs)"],
                        "commentary_ref_path": COMMENTARY_PATH},
        "counts": {"sutras": len(records), "extras": sorted(extras),
                   "conflicts": {k: len(v) for k, v in conflicts.items()}},
        "credit": "Sūtra map workbook (paanini-ashtadhyaayi-sutra-map) and ashtadhyayi.com "
                  "(github.com/ashtadhyayi-com/data); imported read-only, provenance per field.",
    }), encoding="utf-8")
    print(f"{len(records)} sūtras, {len(extras)} extras, overrides {applied}/{n_over}, "
          + ", ".join(f"{k} conflicts {len(v)}" for k, v in conflicts.items()))


if __name__ == "__main__":
    main()
